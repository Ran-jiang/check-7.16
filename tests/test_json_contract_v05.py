"""Claim 0.6 JSON 从识别快照恢复后可以直接完成核验。"""

from docx import Document

from ccitecheck.application import (
    extract_document_claims,
    parse_and_validate_document,
    verify_document_claims,
)
from ccitecheck.domain.citation import ClaimDocument
from ccitecheck.domain.checks import CheckVerdict
from ccitecheck.domain.statute_results import StatuteMeaningCheck
from ccitecheck.infrastructure.database import connect, init_db, upsert_article, upsert_law


def test_new_claim_json_roundtrip_runs_verification(tmp_path):
    law_db = tmp_path / "laws.sqlite"
    init_db(law_db)
    with connect(law_db) as connection:
        law_id = upsert_law(connection, {
            "title": "中华人民共和国民法典",
            "source_type": "law",
        })
        upsert_article(connection, law_id, {
            "article_no": "第五百七十七条",
            "text": "当事人一方不履行合同义务或者履行合同义务不符合约定的，应当承担违约责任。",
        })

    document_path = tmp_path / "contract.docx"
    document = Document()
    document.add_paragraph(
        "依据《中华人民共和国民法典》第五百七十七条，被告应当承担违约责任。"
    )
    document.save(document_path)

    parsed = parse_and_validate_document(document_path)
    recognized = extract_document_claims(parsed)
    payload = recognized.model_dump_json()
    restored = ClaimDocument.model_validate_json(payload)
    verified = verify_document_claims(restored, law_db, semantic_check=False)

    assert restored.claim_meta.schema_version == "0.6"
    assert restored.claim_meta.extractor_version == "0.5"
    source = restored.claims[0].entities.legal_sources[0]
    assert source.recognition.resolver == "direct"
    assert not hasattr(restored.claims[0].entities.citations[0], "verification")
    assert verified.schema_version == "0.9"
    assert verified.statute_results[0].claim_text == restored.claims[0].text
    assert not hasattr(verified.statute_results[0], "document_quote")
    assert verified.statute_results[0].law_identity_resolver == "direct"
    assert verified.statute_results[0].evidence.article_no == "第五百七十七条"


def test_surface_title_citation_uses_canonical_title_for_lookup(tmp_path):
    law_db = tmp_path / "laws.sqlite"
    init_db(law_db)
    with connect(law_db) as connection:
        law_id = upsert_law(connection, {
            "title": "中华人民共和国民法典",
            "source_type": "law",
        })
        upsert_article(connection, law_id, {
            "article_no": "第一条",
            "text": "为了保护民事主体的合法权益，制定本法。",
        })

    payload = {
        "claim_meta": {
            "schema_version": "0.6",
            "extractor_version": "0.5",
        },
        "claims": [{
            "claim_id": "cl_00001",
            "claim_type": "legal_source_claim",
            "text": "依据《民法典》第一条，应依法处理。",
            "anchor_ids": ["line00001"],
            "entities": {
                "legal_sources": [{
                    "title": "民法典",
                    "canonical_title": "中华人民共和国民法典",
                    "articles": [{"article": "第一条"}],
                }],
                "citations": [{
                    "law_title": "民法典",
                    "locator": {"article": "第一条"},
                    "citation_span": [2, 10],
                    "span_status": "located",
                }],
            },
        }],
    }
    restored = ClaimDocument.model_validate(payload)
    verified = verify_document_claims(restored, law_db, semantic_check=False)

    assert verified.statute_results[0].evidence.law_title == "中华人民共和国民法典"


def test_semantic_comparison_receives_complete_claim_text(tmp_path):
    law_db = tmp_path / "laws.sqlite"
    init_db(law_db)
    with connect(law_db) as connection:
        law_id = upsert_law(connection, {
            "title": "中华人民共和国民法典",
            "source_type": "law",
        })
        upsert_article(connection, law_id, {
            "article_no": "第五百七十七条",
            "text": "当事人一方不履行合同义务或者履行合同义务不符合约定的，应当承担违约责任。",
        })

    text = "对被告的行为应当适用《中华人民共和国民法典》第五百七十七条的规定。"
    document_path = tmp_path / "full-claim.docx"
    document = Document()
    document.add_paragraph(text)
    document.save(document_path)

    class CaptureChecker:
        def __init__(self):
            self.claim_text = None

        def compare(self, claim_text, cited_source, evidence):
            self.claim_text = claim_text
            return StatuteMeaningCheck(verdict=CheckVerdict.PASS)

    checker = CaptureChecker()
    recognized = extract_document_claims(parse_and_validate_document(document_path))
    from ccitecheck.application.verify_claims import verify_claim_document

    verified = verify_claim_document(
        recognized, law_db, semantic_checker=checker, include_cases=False
    )

    assert checker.claim_text == text
    assert verified.statute_results[0].claim_text == text
    assert "verification" not in recognized.model_dump_json()
