from ccitecheck.recognition.cases import extract_case_refs
from ccitecheck.recognition.statutes import extract_legal_sources


def test_complete_judgment_title_without_case_number_is_recognized():
    text = (
        "李某、河南金钢门业股份有限公司统一社会信用代码等"
        "违反安全保障义务责任纠纷民事再审民事判决书"
    )

    refs = extract_case_refs(text)

    assert len(refs) == 1
    assert refs[0].case_name == text
    assert refs[0].case_number is None
    assert refs[0].document_type == "判决书"


def test_plain_judgment_document_description_is_not_a_case_title():
    assert extract_case_refs("当事人向法院提交了民事判决书作为证据。") == []


def test_compact_arabic_article_range_is_expanded():
    sources = extract_legal_sources("依据《消防法》58-69条承担相应责任。")

    assert [item.article for item in sources[0].articles] == [
        f"第{number}条" for number in range(58, 70)
    ]


def test_document_number_range_is_not_an_article_range():
    sources = extract_legal_sources("依据《消防法》及国发〔2018-2020〕号文件处理。")

    assert sources[0].articles == []


def test_court_boundary_excludes_grammar_prefix():
    refs = extract_case_refs(
        "申请人不服山西省高级人民法院（2025）晋民再174号民事判决。"
    )

    assert refs[0].court == "山西省高级人民法院"


def test_nearest_court_alias_is_normalized_for_each_case():
    refs = extract_case_refs(
        "（2021）最高法民申1516号案中，最高人民法院作出认定；"
        "（2015）潍商终字第877号案中潍坊中院作出认定。"
    )

    assert refs[0].court == "最高人民法院"
    assert refs[1].court == "潍坊市中级人民法院"


def test_short_name_declaration_allows_document_number_metadata():
    from ccitecheck.recognition.statutes import extract_alias_declarations

    text = (
        "《山西省农村信用社会计业务印章管理办法》"
        "（晋农信发〔2014〕28号，下称《印章管理办法》）"
        "第9条、第11条明确规定。"
    )
    sources = extract_legal_sources(
        text
    )

    assert [source.title for source in sources] == [
        "山西省农村信用社会计业务印章管理办法", "印章管理办法"
    ]
    assert [item.article for item in sources[1].articles] == ["第9条", "第11条"]
    assert extract_alias_declarations(text)[0].alias_raw == "印章管理办法"


def test_case_direct_quote_trigger_extracts_quote_text():
    from ccitecheck.recognition.cases import find_holding_trigger_position

    text = "（2021）最高法民申1516号民事裁定书载明：‘业务专用章具有对外效力。’"
    refs = extract_case_refs(text)
    start = find_holding_trigger_position(text, refs)

    assert start is not None
    assert text[start:] == "‘业务专用章具有对外效力。’"


def test_docx_body_reference_is_linked_to_footnote_block(tmp_path):
    import zipfile

    from docx import Document

    from ccitecheck.domain.document import BlockRelationType, BlockType
    from ccitecheck.parsing.docx import parse_docx

    source_path = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("正文说明。")
    document.save(source_path)

    footnotes_xml = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
    <w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:footnote w:id="2"><w:p><w:r><w:t>依据《民法典》第五百零九条规定，当事人应当全面履行义务。</w:t></w:r></w:p></w:footnote>
    </w:footnotes>'''.encode("utf-8")
    rebuilt = tmp_path / "with-reference.docx"
    with zipfile.ZipFile(source_path) as source, zipfile.ZipFile(rebuilt, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "word/document.xml":
                content = content.replace(
                    b"</w:r></w:p>",
                    b'</w:r><w:r><w:footnoteReference w:id="2"/></w:r></w:p>',
                    1,
                )
            target.writestr(item, content)
        target.writestr("word/footnotes.xml", footnotes_xml)

    parsed = parse_docx(str(rebuilt))
    body = next(block for block in parsed.blocks if block.type == BlockType.PARAGRAPH)
    note = next(block for block in parsed.blocks if block.type == BlockType.FOOTNOTE)

    assert body.note_references[0].note_id == "2"
    assert body.note_references[0].char_offset == len("正文说明。")
    assert any(
        relation.relation_type == BlockRelationType.NOTE_REFERENCE
        and relation.target_block_id == note.block_id
        for relation in body.relations
    )

    from ccitecheck.parsing import build_chunks
    from ccitecheck.recognition import extract_claims

    claim_document = extract_claims(build_chunks(parsed), include_cases=False)
    note_claim = next(claim for claim in claim_document.claims if claim.note_context)
    assert note_claim.note_context.note_id == "2"
    assert note_claim.note_context.referenced_from[0].block_id == "word:p:0"
    assert note_claim.text == "依据《民法典》第五百零九条规定，当事人应当全面履行义务。"
    assert not hasattr(note_claim.entities.citations[0], "verification")


def test_statute_citation_keeps_claim_text_as_single_comparison_source():
    from ccitecheck.domain.citation import (
        Claim,
        ClaimType,
        LegalSourceClaimEntities,
    )
    from ccitecheck.recognition.spans import locate_claim_article_spans

    text = "依据《民法典》第五百零九条规定，当事人应当按照约定全面履行自己的义务。"
    claim = Claim(
        claim_id="cl_00001",
        claim_type=ClaimType.LEGAL_SOURCE_CLAIM,
        text=text,
        anchor_ids=["line00001"],
        entities=LegalSourceClaimEntities(legal_sources=extract_legal_sources(text)),
    )

    locate_claim_article_spans(claim)
    assert claim.text == text
    assert not hasattr(claim.entities.citations[0], "verification")


def test_quote_before_statute_citation_keeps_complete_claim_text():
    from ccitecheck.domain.citation import Claim, ClaimType, LegalSourceClaimEntities
    from ccitecheck.recognition.spans import locate_claim_article_spans

    text = "“当事人应当按照约定全面履行自己的义务。”——《民法典》第五百零九条。"
    claim = Claim(
        claim_id="cl_00001",
        claim_type=ClaimType.LEGAL_SOURCE_CLAIM,
        text=text,
        anchor_ids=["line00001"],
        entities=LegalSourceClaimEntities(legal_sources=extract_legal_sources(text)),
    )

    locate_claim_article_spans(claim)
    assert claim.text == text
    assert not hasattr(claim.entities.citations[0], "verification")


def test_nested_authoritative_text_does_not_replace_claim_text():
    from ccitecheck.orchestration.scheduler import _CheckItem
    from ccitecheck.domain.citation import (
        ArticleRef,
        Claim,
        ClaimType,
        LegalSourceClaimEntities,
    )
    article = ArticleRef(article="第二十五条")
    claim = Claim(
        claim_id="cl_00001",
        claim_type=ClaimType.LEGAL_SOURCE_CLAIM,
        text="国家有关规定包括《借贷规定》第二十五条。",
        anchor_ids=["line00001"],
        entities=LegalSourceClaimEntities(),
    )
    item = _CheckItem(
        claim=claim,
        law_title="借贷规定",
        display_title="借贷规定",
        article=article,
        article_no="第二十五条",
        not_verifiable=None,
        relation_parent_authoritative_text="北大法宝返回的父法条权威原文",
    )

    assert item.claim.text == "国家有关规定包括《借贷规定》第二十五条。"


def test_bare_law_identity_is_kept_as_raw_candidate():
    from ccitecheck.domain.document import Anchor, Block, BlockType, ParsedDocument
    from ccitecheck.recognition.arbitration import arbitrate_claim_candidates
    from ccitecheck.recognition.rules import extract_rule_candidates
    from ccitecheck.recognition.service import build_indexes

    text = "依照星河数据治理法第十条处理。"
    block = Block(
        block_id="b_00001",
        type=BlockType.PARAGRAPH,
        text=text,
        body_order=0,
        block_order=0,
        para_index=0,
        sentence_anchors=["line00001"],
        anchor_range=["line00001", "line00001"],
    )
    document = ParsedDocument(
        blocks=[block],
        anchors=[Anchor(
            anchor="line00001",
            text=text,
            block_id=block.block_id,
            para_index=0,
            char_start=0,
            char_end=len(text),
        )],
    )
    candidates = extract_rule_candidates(document, build_indexes(document))
    claims = arbitrate_claim_candidates(candidates, document)

    assert len(claims) == 1
    source = claims[0].entities.legal_sources[0]
    assert source.title == "星河数据治理法"
    assert source.raw_title_candidate == "星河数据治理法"
    assert source.canonical_title is None
    assert source.articles[0].article == "第十条"


def test_quote_and_following_statute_citation_across_paragraphs_share_claim(tmp_path):
    from docx import Document

    from ccitecheck.application.check_document import (
        extract_document_claims,
        parse_and_validate_document,
    )
    from ccitecheck.recognition.spans import locate_claim_article_spans

    path = tmp_path / "quote-before-citation.docx"
    document = Document()
    document.add_paragraph("“当事人应当按照约定全面履行自己的义务。”")
    document.add_paragraph("——《中华人民共和国民法典》第五百零九条。")
    document.save(path)

    parsed = parse_and_validate_document(path)
    claims = extract_document_claims(parsed)
    claim = next(item for item in claims.claims if item.claim_type.value == "legal_source_claim")
    locate_claim_article_spans(claim)

    assert len(claim.anchor_ids) == 2
    assert claim.text == "“当事人应当按照约定全面履行自己的义务。”——《中华人民共和国民法典》第五百零九条。"


def test_each_carry_forward_paragraph_keeps_its_own_citation_occurrence():
    from ccitecheck.domain.citation import Claim, ClaimType, LegalSourceClaimEntities
    from ccitecheck.recognition.spans import locate_claim_article_spans

    text = "《反不正当竞争法》第十三条第三款规定A。同条第四款规定B。第五款进一步规定C。"
    claim = Claim(
        claim_id="cl_00001",
        claim_type=ClaimType.LEGAL_SOURCE_CLAIM,
        text=text,
        anchor_ids=["line00001"],
        entities=LegalSourceClaimEntities(legal_sources=extract_legal_sources(text)),
    )

    locate_claim_article_spans(claim)

    assert [item.locator.paragraph for item in claim.entities.citations] == [
        "第三款", "第四款", "第五款"
    ]
    assert [item.role for item in claim.entities.citations] == [
        "direct", "carry_forward", "carry_forward"
    ]


def test_recognition_pipeline_keeps_successive_carry_forward_sentences(tmp_path):
    from docx import Document

    from ccitecheck.application.check_document import (
        extract_document_claims,
        parse_and_validate_document,
    )

    path = tmp_path / "carry-forward.docx"
    document = Document()
    document.add_paragraph(
        "《中华人民共和国反不正当竞争法》第十三条第三款规定A。"
        "同条第四款规定B。第五款进一步规定C。"
    )
    document.save(path)

    claim_document = extract_document_claims(parse_and_validate_document(path))
    statute_claims = [
        claim for claim in claim_document.claims
        if claim.claim_type.value == "legal_source_claim"
    ]

    assert [claim.entities.citations[0].locator.paragraph for claim in statute_claims] == [
        "第三款", "第四款", "第五款",
    ]
    assert [claim.text for claim in statute_claims] == [
        "《中华人民共和国反不正当竞争法》第十三条第三款规定A。",
        "同条第四款规定B。",
        "第五款进一步规定C。",
    ]


def test_unresolved_bare_law_does_not_also_inherit_previous_law(tmp_path):
    from docx import Document

    from ccitecheck.application.check_document import (
        extract_document_claims,
        parse_and_validate_document,
    )

    path = tmp_path / "unresolved-after-explicit.docx"
    document = Document()
    document.add_paragraph(
        "《中华人民共和国民法典》第五百零九条规定A。"
        "依照星河数据治理法第十条处理。"
    )
    document.save(path)

    claim_document = extract_document_claims(parse_and_validate_document(path))
    raw_candidate = next(
        claim for claim in claim_document.claims
        if any(
            source.raw_title_candidate == "星河数据治理法"
            for source in claim.entities.legal_sources
        )
    )

    assert [source.title for source in raw_candidate.entities.legal_sources] == [
        "星河数据治理法"
    ]


def test_closed_quoted_term_does_not_break_nested_relation_candidate():
    from ccitecheck.recognition.relations import relation_candidates
    from ccitecheck.orchestration.scheduler import _collect_check_items
    from ccitecheck.domain.citation import Claim, ClaimType, LegalSourceClaimEntities
    from ccitecheck.recognition.spans import locate_claim_article_spans

    text = (
        "《中华人民共和国民法典》第六百八十条规定，禁止高利放贷；"
        "其中“国家有关规定”包括"
        "《最高人民法院关于审理民间借贷案件适用法律若干问题的规定》第二十五条。"
    )
    claim = Claim(
        claim_id="cl_nested",
        claim_type=ClaimType.LEGAL_SOURCE_CLAIM,
        text=text,
        anchor_ids=["line00001"],
        entities=LegalSourceClaimEntities(legal_sources=extract_legal_sources(text)),
    )
    locate_claim_article_spans(claim)
    items = _collect_check_items(type("ClaimDoc", (), {"claims": [claim]})())

    assert relation_candidates(items) == {1: [0]}
    assert all(not hasattr(item, "verification") for item in claim.entities.citations)
