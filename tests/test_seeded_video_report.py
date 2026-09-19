from docx import Document

from ccitecheck.application import extract_document_claims, parse_and_validate_document
from ccitecheck.domain.citation import ClaimType
from ccitecheck.orchestration.scheduler import _collect_check_items
from ccitecheck.recognition.statutes import extract_alias_declarations, extract_legal_sources
from ccitecheck.retrieval.sources.pkulaw.parsing import parse_semantic_case_text
from ccitecheck.verification.semantic import (
    _application_check_from_raw,
    _case_reasoning_check_from_raw,
)
from ccitecheck.verification.cases.identity import equivalent_case_name


def test_seeded_report_recognizes_aliases_foreign_sources_and_case_holding(tmp_path):
    lines = [
        "《生成式人工智能服务管理暂行办法》（以下简称“《暂行办法》”）第八条规定训练数据应当来源合法。",
        "《暂行办法》第二条明确内部研发不适用该办法。据此，内部研发可以一并豁免其他法定义务。",
        "最高人民法院发布的典型案例中，包括彭某某诉某软件运营公司肖像权纠纷案。涉案公司提供AI换脸模板。法院判令赔偿3000元。该案确立的规则是，营利使用肖像即构成侵权。因此获利是侵权的充分依据。",
        "欧盟《人工智能法》（Regulation (EU) 2024/1689，以下简称“EU AI Act”）第50条第2款要求部署者标记输出。",
        "EU AI Act第50条第4款允许虚构作品免于披露。",
        "美国《TAKE IT DOWN Act》第3条要求72小时内删除。",
    ]
    path = tmp_path / "seeded.docx"
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    document.save(path)

    claims = extract_document_claims(parse_and_validate_document(path))
    case = next(claim for claim in claims.claims if claim.claim_type == ClaimType.CASE_HOLDING_PARAPHRASE)
    assert [ref.case_name for ref in case.entities.case_refs] == ["彭某某诉某软件运营公司肖像权纠纷案"]
    assert "充分依据" in case.text
    internal = next(claim for claim in claims.claims if "一并豁免" in claim.text)
    assert "不适用该办法" in internal.text

    items = _collect_check_items(claims, "data/laws.sqlite")
    assert [(item.display_title, item.law_title, item.jurisdiction) for item in items] == [
        ("暂行办法", "生成式人工智能服务管理暂行办法", "CN"),
        ("暂行办法", "生成式人工智能服务管理暂行办法", "CN"),
        ("人工智能法", "人工智能法", "EU"),
        ("EU AI Act", "EU AI Act", "EU"),
        ("TAKE IT DOWN Act", "TAKE IT DOWN Act", "US"),
    ]


def test_seeded_alias_and_plain_act_shapes():
    chinese = "《人工智能生成合成内容标识办法》（以下简称“《标识办法》”）第九条"
    eu = "欧盟《人工智能法》（Regulation (EU) 2024/1689，以下简称“EU AI Act”）第50条第2款"
    assert extract_alias_declarations(chinese)[0].alias_raw == "标识办法"
    assert extract_alias_declarations(eu)[0].alias_raw == "EU AI Act"
    source = extract_legal_sources("EU AI Act第50条第4款")[0]
    assert source.title == "EU AI Act"
    assert source.articles[0].paragraphs == ["第4款"]


def test_pkulaw_semantic_case_text_is_evidence_not_source_error():
    records = parse_semantic_case_text(
        "共返回 1 条案例：\n\n"
        "1. [典型案例] 最高人民法院典型案例之三：彭某某诉某软件运营公司肖像权纠纷案——擅用肖像换脸应担责 | \n"
        "   审理法院：最高人民法院 | 审结日期：2025-06-12\n"
        "   本院认为：该公司未经彭某某授权同意，以营利为目的使用其肖像。\n"
        "   裁判结果：赔偿3000元。\n"
        "   链接：https://pkulaw.com/pfnl/example.html\n"
    )
    assert len(records) == 1
    assert "彭某某诉某软件运营公司肖像权纠纷案" in records[0].title
    assert "未经彭某某授权同意" in records[0].holding
    assert records[0].url == "https://pkulaw.com/pfnl/example.html"
    assert records[0].case_number == ""


def test_case_bibliographic_assertion_is_not_a_holding_error():
    result = _case_reasoning_check_from_raw(
        {
            "verdict": "pass",
            "assertions": [{"id": 1, "judgment": "not_holding", "hit_sentence_ids": []}],
        },
        ["最高人民法院于2025年6月12日发布该典型案例。"],
        ["该公司未经授权同意，以营利为目的使用肖像。"],
        False,
    )
    assert result.findings == []


def test_seeded_case_name_matches_only_decorated_typical_case_title():
    cited = "彭某某诉某软件运营公司肖像权纠纷案"
    decorated = "最高人民法院发布6个利用网络、信息技术侵害人格权典型案例之三：彭某某诉某软件运营公司肖像权纠纷案——擅用肖像应担责"
    assert equivalent_case_name(cited, decorated)
    assert not equivalent_case_name("甲公司诉乙公司合同纠纷案", "甲公司诉乙公司合同纠纷案再审审查案")


def test_overlong_application_summary_is_safely_bounded():
    result = _application_check_from_raw({
        "verdict": "review",
        "comparison": "摘" * 500,
        "reviews": [{
            "error_type": "rule_fact_mismatch",
            "summary": "错" * 500,
            "suggestion": "修改",
            "related_sources": [],
        }],
    })
    assert len(result.comparison) == 300
    assert len(result.reviews[0].summary) == 300
