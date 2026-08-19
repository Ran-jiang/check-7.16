from docx import Document

from apps.api.docx_text import sanitize_for_docx

# python-docx 拒绝的全部码位，与 sanitize_for_docx 的替换范围一一对应。
XML_ILLEGAL_CODES = [*range(0x00, 0x09), 0x0B, 0x0C, *range(0x0E, 0x20)]


def test_replaces_every_illegal_code_point_without_changing_length():
    for code in XML_ILLEGAL_CODES:
        line = f"第61条{chr(code)}及第108条"
        cleaned = sanitize_for_docx(line)
        assert len(cleaned) == len(line)
        assert cleaned == "第61条 及第108条"
        Document().add_paragraph(cleaned)


def test_keeps_xml_legal_whitespace_untouched():
    line = "第61条\t及第108条\r\n结尾"
    assert sanitize_for_docx(line) == line


def test_footnote_marks_no_longer_break_docx_construction():
    """选中带脚注的条文曾使 /api/checks/selection 抛 ValueError 返回 500。"""
    line = "依据《中华人民共和国民法典》第61条\x02及第108条\x02，李斌承担责任"
    Document().add_paragraph(sanitize_for_docx(line))
